from bs4 import BeautifulSoup
import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill
from selenium import webdriver
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
import os
from openpyxl.styles import NamedStyle


def listSeen():
    options = webdriver.ChromeOptions()
    options.add_argument("--lang=ja")  # <-- Forzar idioma japonés

    driver = webdriver.Chrome(
        service=Service(ChromeDriverManager().install()),
        options=options
    )
    driver.get("https://myanimelist.net/animelist/Walter135?order=3&order2=-14&status=7")
    try:
        elem = driver.find_element(By.TAG_NAME, "body")
        driver.find_element(By.TAG_NAME, "td")
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-53393"))  # Tengoku Daimakyou
        )
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-42361"))  # Ijiranaide, Nagatoro-san
        )
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-4224"))  # Toradora
        )
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-11111"))  # Another
        )
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-31240"))  # Re:Zero
        )
        elem.send_keys(Keys.END)

        element = WebDriverWait(driver, 5).until(
            EC.presence_of_element_located((By.ID, "more-34618"))  # Blend S
        )
        elem.send_keys(Keys.END)

        soup = BeautifulSoup(driver.page_source, 'html.parser')
        driver.quit()
        divGeneral = soup.find_all("tbody", class_="list-item")

        for divs in divGeneral:
            nameTitle = divs.find_all('a', class_='link sort')[1].text
            flag = False
            for l in listNoProcces:
                if l in nameTitle:
                    flag = True

            if flag:
                continue

            cantW = len(divs.find_all('td', class_='watching'))
            cantO = len(divs.find_all('td', class_='onhold'))
            cantP = len(divs.find_all('td', class_='plantowatch'))

            for j in letterProhibited:
                nameTitle = nameTitle.replace(j, "")
            nameList.append(nameTitle)
            urlList.append(base + divs.find_all('a', class_='link sort')[1]['href'])
            startList.append(divs.find_all('td', class_='started')[0].text)

            if cantW > 0:
                typeList.append("Watching")
            else:
                if cantO > 0:
                    typeList.append("On hold")
                else:
                    if cantP > 0:
                        typeList.append("Plan to watch")
                    else:
                        typeList.append("Completed")
    except:
        driver.close()
        driver.quit()
        listSeen()


def Excel(title):
    df = pd.DataFrame({"Nombre": nameList, "Fecha": startList, "Tipo": typeList, "Url": urlList})

    df.to_excel(title, index=False)

    wb = openpyxl.load_workbook(filename=title)
    worksheet = wb.active
    worksheet.auto_filter.ref = worksheet.dimensions
    date_style = NamedStyle(name='datetime', number_format='DD/MM/YYYY')

    for col in worksheet.columns:
        max_length = 0
        column = col[0].column_letter  # Get the column name
        k = 1
        for cell in col:
            if column == "A" and k != 1:
                worksheet[column + str(k)].hyperlink = worksheet["D" + str(k)].value
            if (column == "B" or column == "C") and k != 1:
                worksheet[column + str(k)].style = date_style
            try:  # Necessary to avoid error on empty cells
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
            k = k + 1
        adjusted_width = (max_length + 2) * 1.2
        worksheet.column_dimensions[column].width = adjusted_width

    blueFill = PatternFill(start_color='05A9B1',
                           end_color='05A9B1',
                           fill_type='solid')
    greenFill = PatternFill(start_color='2db039',
                            end_color='2db039',
                            fill_type='solid')
    yellowFill = PatternFill(start_color='f1c83e',
                             end_color='f1c83e',
                             fill_type='solid')
    grayFill = PatternFill(start_color='c3c3c3',
                           end_color='c3c3c3',
                           fill_type='solid')

    row = 1
    for cell in worksheet['C']:
        if cell.value == "Watching":
            worksheet['A' + str(row)].fill = greenFill
        if cell.value == "Completed":
            worksheet['A' + str(row)].fill = blueFill
        if cell.value == "On hold":
            worksheet['A' + str(row)].fill = yellowFill
        if cell.value == "Plan to watch":
            worksheet['A' + str(row)].fill = grayFill
        row = row + 1

    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.delete_cols(4, 1)
    wb.save(title)
    os.system(title)


listNoProcces = ["One Piece"]
letterProhibited = ["\\", "/", ":", "*", "?", "<", ">", "|", '"']
base = "https://myanimelist.net"
nameList = list()
urlList = list()
startList = list()
typeList = list()


#listSeen()
#Excel("Excel\\ListSeen.xlsx")
